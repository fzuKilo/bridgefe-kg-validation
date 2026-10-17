"""导出补充材料 S1：各方案的提示词与工具定义（论文 4.7 节）。

  python -m kgfem.export_prompts --out paper_supplementary
生成 Markdown（S1_prompts.md）；装有 python-docx 时另生成 Word 版（S1_prompts.docx）。
"""
import argparse
import json
import os

from .cli import ROOT


def sections():
    from .agent.baselines import SELF_CHECK, knowledge_text, system_prompt
    from .agent.claude import SYSTEM_PROMPT
    from .agent.direct import C0_SYSTEM, C1_SYSTEM
    from .agent.tools import DEFECT_RULE, DEFECT_TOOL, SESSION_TOOLS
    from .faults import REPAIR_PROMPT
    out = [('C5 — KG-constrained agent (proposed): system prompt', SYSTEM_PROMPT),
           ('C0 — direct script generation: system prompt', C0_SYSTEM),
           ('C1 — direct decision generation: system prompt', C1_SYSTEM)]
    for cfg, title in (('C2', 'agent with prompt knowledge'), ('C3', 'agent with retrieval augmentation'),
                       ('C4', 'agent with the KG as tools, without validation')):
        p = system_prompt(cfg)
        if cfg == 'C2':                      # 规则与条文全文很长，单列一节
            p = p.replace(knowledge_text(), '<rules and clauses of the KG, listed in the next section>')
        out.append((f'{cfg} — {title}: system prompt', p))
        if cfg == 'C2':
            out.append(('Knowledge text written into the C2 prompt', knowledge_text()))
    out.append(('C2 — self-check step', SELF_CHECK))
    out.append(('Fault-injection experiment — repair prompt (Section 5.5.2)', REPAIR_PROMPT))
    out.append(('Boundary-input experiment — additional prompt paragraph (Section 5.6)', DEFECT_RULE.strip()))
    tools = [{k: t[k] for k in ('name', 'description', 'input_schema')} for t in SESSION_TOOLS + [DEFECT_TOOL]]
    out.append(('Tool definitions (13 KG and modeling tools; report_input_defect only in the boundary-input '
                'experiment)', json.dumps(tools, ensure_ascii=False, indent=2)))
    return [(f'S1.{i} {t}', b) for i, (t, b) in enumerate(out, 1)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default=os.path.join(ROOT, 'paper_supplementary'))
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    secs = sections()
    title = 'Supplementary Material S1. Prompts and tool definitions'
    intro = ('Knowledge graph-constrained large language model agent for automated finite element modeling of girder '
             'bridges. The texts below are exported verbatim from the implementation used in the experiments.')
    md = [f'# {title}', '', intro, '']
    for t, body in secs:
        md += [f'## {t}', '', '```', body.strip(), '```', '']
    path = os.path.join(a.out, 'S1_prompts.md')
    open(path, 'w', encoding='utf-8').write('\n'.join(md))
    print(path)
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(intro)
    for t, body in secs:
        doc.add_heading(t, level=2)
        for line in body.strip().split('\n'):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.name = 'Courier New'
            r.font.size = Pt(8)
            p.paragraph_format.space_after = Pt(0)
    path = os.path.join(a.out, 'S1_prompts.docx')
    doc.save(path)
    print(path)


if __name__ == '__main__':
    main()
